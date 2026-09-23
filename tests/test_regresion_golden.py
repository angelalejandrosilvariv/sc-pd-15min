"""Regresión byte/valor de la entrega contractual y sus fórmulas.

La referencia fue capturada antes de habilitar el caché de lectura. El XLSX de
entrada es la referencia sintética exhaustiva de ``test_generar_entrega_cen``.
"""
import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path[:0] = [str(Path(__file__).parent), str(Path(__file__).parents[1] / "scripts")]
from generar_entrega_cen import generar_entrega  # noqa: E402
from test_generar_entrega_cen import reporte as fixture_reporte  # noqa: E402

DISEÑO_PANEL = [
    {},
    {"RESOLUCION_MARGEN": "hora"},
    {"TARIFA_CONFIGURACION": "instruida"},
    {"PARTIDA_EN_PRUEBAS": "validar_orden_om_fallida"},
    {"PARTIDA_EN_PRUEBAS": "validar_si_queda_disponible_om"},
    {"HORAS_SIN_HISTORIA": "nulo"},
    {"DIFERIR_CICLOS_SIN_TERMINAR": 0},
    {"RUTA_REPORTE_MES_PASADO": ""},
]


def _formulas(libro):
    wb = load_workbook(libro, data_only=False, read_only=True)
    return {ws.title: {c.coordinate: c.value for fila in ws.iter_rows() for c in fila
                       if c.data_type == "f"} for ws in wb.worksheets
            if any(c.data_type == "f" for fila in ws.iter_rows() for c in fila)}


@pytest.mark.parametrize("panel", DISEÑO_PANEL, ids=lambda p: next(iter(p), "defaults"))
def test_entrega_y_formulas_identicas_a_golden(tmp_path, panel):
    reporte = fixture_reporte.__wrapped__(tmp_path)
    carpeta = generar_entrega(reporte, tmp_path / "resultado", panel=panel)
    with (Path(__file__).parent / "golden" / "entrega_sintetica.json").open("r", encoding="utf8") as f:
        golden = json.load(f)
    actuales = {p.name.split("_", 2)[-1]: p.read_text(encoding="utf-8-sig")
                for p in carpeta.glob("*.csv")}
    # parámetros contiene fecha, commit, hash del XLSX y el panel deliberadamente
    # variable; se valida en las pruebas de entrega, no como artefacto inmutable.
    actuales.pop("parametros.csv")
    esperados = dict(golden["csv"])
    esperados.pop("parametros.csv")
    assert actuales == esperados
    formulas = _formulas(next(carpeta.glob("*.xlsx")))
    if not panel:
        assert {k: v for k, v in formulas.items() if v} == golden["formulas"]
    else:
        # Algunas opciones cambian deliberadamente la plantilla contractual
        # (p. ej. margen horario), pero nunca deben materializarla como valores.
        assert sum(map(len, formulas.values())) >= 200
