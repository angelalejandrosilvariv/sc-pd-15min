"""Equivalencia de la exportación rápida del motor con DataFrame.to_excel."""
from datetime import date
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from pandas.testing import assert_frame_equal

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from sc_pd_motor_v7 import (crear_formato_encabezado_xlsxwriter,
                            escribir_hoja_xlsxwriter)


def _ajustar_anchos(ws, df):
    for idx, columna in enumerate(df.columns):
        largo = df.iloc[:, idx].map(lambda valor: len(str(valor))).max() if len(df) else 0
        ws.set_column(idx, idx, min(max(int(largo or 0), len(str(columna))) + 2, 50))


def test_exportacion_rapida_equivale_a_to_excel_en_valores_tipos_y_anchos(tmp_path):
    df = pd.DataFrame({
        "texto": pd.Series(["", "abc", pd.NA, "más largo"], dtype="string"),
        "objeto": [None, "x", "", pd.NA],
        "booleano": [True, False, True, False],
        "entero": pd.Series([1, 0, -7, 42], dtype="int64"),
        "flotante": [1.5, float("nan"), -2.25, 0.0],
        "fecha_hora": pd.to_datetime([
            "2026-01-02 03:04:05", None, "2026-12-31 23:59:00", "2026-06-01"], format="mixed"),
        "fecha_hora_cero": pd.to_datetime([
            "2026-01-02", None, "2026-12-31", pd.NaT], format="mixed"),
        "fecha": [date(2026, 1, 2), None, date(2026, 12, 31), date(2026, 6, 1)],
    })
    referencia = tmp_path / "pandas.xlsx"
    rapido = tmp_path / "rapido.xlsx"

    with pd.ExcelWriter(referencia, engine="xlsxwriter") as writer:
        df.to_excel(writer, sheet_name="Datos", index=False)
        _ajustar_anchos(writer.sheets["Datos"], df)

    with pd.ExcelWriter(
            rapido, engine="xlsxwriter",
            engine_kwargs={"options": {"constant_memory": True}}) as writer:
        encabezado = crear_formato_encabezado_xlsxwriter(writer, df)
        fecha = writer.book.add_format({"num_format": writer.date_format})
        fecha_hora = writer.book.add_format({"num_format": writer.datetime_format})
        escribir_hoja_xlsxwriter(
            writer, "Datos", df, encabezado, fecha, fecha_hora)

    esperado = pd.read_excel(referencia, sheet_name="Datos")
    obtenido = pd.read_excel(rapido, sheet_name="Datos")
    assert_frame_equal(obtenido, esperado, check_exact=True)

    wb_referencia = load_workbook(referencia)
    wb_rapido = load_workbook(rapido)
    ws_referencia = wb_referencia["Datos"]
    ws_rapido = wb_rapido["Datos"]
    columnas = [celda.column_letter for celda in ws_referencia[1]]
    assert [ws_rapido.column_dimensions[c].width for c in columnas] == [
        ws_referencia.column_dimensions[c].width for c in columnas]
    assert ws_rapido["A1"]._style == ws_referencia["A1"]._style
    assert ws_rapido["F2"].number_format == ws_referencia["F2"].number_format
    assert ws_rapido["G2"].number_format == ws_referencia["G2"].number_format
    assert ws_rapido["H2"].number_format == ws_referencia["H2"].number_format
