"""Pruebas sintéticas de la entrega con formato del Excel horario (spec 32)."""
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parents[1] / "scripts"))
from generar_entrega_cen import (  # noqa: E402
    CICLO, FORMULAS, HOJA_PAGOS, PAGOS, PD, SHEETS, XHYC, generar_entrega)

T0 = pd.Timestamp("2026-08-01 00:00")
Q = pd.Timedelta(minutes=15)
RIO_C_PARTIDA = pd.Timestamp("2026-07-31 23:50")
RIO_C_DETENCION = pd.Timestamp("2026-08-01 00:20")


def _bloque(etiqueta, rel, config, t, gen=1.0, tarifa_p=10.0, tarifa_d=1.0, motivo="OM", eo="N",
            consigna="", comentario="", config_rio=None, fuente=None, disponible=1, tipo="Fria", exencion=False):
    return {"Etiqueta_Relacionada": etiqueta, "Central_Relacionada": rel, "Ciclo_ID_Relacionada": int(etiqueta.split("&")[1]),
            "Central": config, "UNIDAD GENERADORA": config, "FECHA_HORA": t, "GENERACION": gen,
            "CMg": 2.0, "CV": 1.0, "Dolar": 1.0, "Valor_Dolar": 1.0, "Margen": gen * 1.0,
            "Costo_Partida": tarifa_p, "Costo_Detencion": tarifa_d, "Costo_Cero": "NO",
            "Filtro_CostoCero_Partida": 1, "Filtro_CostoCero_Detencion": 1, "Tipo_Partida": tipo,
            "Fria_Num1_M": 24, "Tibia_Num2_N": None, "Tibia_Num1_O": 0, "Caliente_Num1_P": 24,
            "MOTIVO": motivo, "ESTADO OPERACIONAL": eo, "CONSIGNAS": consigna, "COMENTARIO": comentario,
            "Configuracion RIO": config_rio or config, "Fuente_Config_RIO": fuente,
            "Disponible (1) / Pruebas (0)": disponible, "Vigencia_RIO": 1, "Filtro_Operacional": 1,
            "Conf despachada RIO": 1, "Llave_FHC": f"260801-1{config}", "Empresa": "E1" if rel != "D" else "E2",
            "Flag_Exencion": exencion}


def _ciclo(etiqueta, rel, estado, inicio, termino, cp, cd, margen, motivo_p="OM", motivo_d="OM",
           coment_p="", coment_d="", exencion=False, fuente_p=None, fuente_d=None, empresa="E1"):
    total = max(0.0, cp + cd - margen)
    return {"Etiqueta_Relacionada": etiqueta, "Central_Relacionada": rel, "Ciclo_ID_Relacionada": int(etiqueta.split("&")[1]),
            "Empresa": empresa, "Ciclo_Mes": "2608", "Estado_Ciclo_Mes": estado, "Inicio_Ciclo": inicio,
            "Termino_Ciclo": termino, "Horas_Detenida_Ciclo": 30.0, "Horas_Cota_Inferior": None,
            "Margen_Suma_Ciclo": margen, "Tipo_Partida": "Fria",
            "Motivo_Partida": motivo_p, "Estado_Op_Partida": "N", "Consigna_Partida": "", "Comentario_Partida": coment_p,
            "Filtro_Conf_Partida": 1, "Filtro_Disp_Partida": 1, "Filtro_Op_Partida": 1 if cp else 0,
            "Vigencia_RIO_Partida": 1, "Filtro_CostoCero_Partida": 1, "Costo_Partida_Base": cp,
            "Costo_Partida_Efectivo": cp,
            "Motivo_Detencion": motivo_d, "Estado_Op_Detencion": "N", "Consigna_Detencion": "", "Comentario_Detencion": coment_d,
            "Filtro_Conf_Detencion": 1, "Filtro_Disp_Detencion": 1, "Filtro_Op_Detencion": 1 if cd else 0,
            "Vigencia_RIO_Detencion": 1, "Filtro_CostoCero_Detencion": 1, "Costo_Detencion_Base": cd,
            "Costo_Detencion_Efectivo": cd, "Costos_Totales_PD": cp + cd, "Total SC_PD": total,
            "Flag_Exencion": exencion, "Fuente_Config_RIO_Partida": fuente_p, "Fuente_Filtros_RIO_Partida": None,
            "Fuente_Config_RIO_Detencion": fuente_d, "Fuente_Filtros_RIO_Detencion": None,
            "Obs_Partida": "Aprobado", "Obs_Detencion": "Aprobado"}


@pytest.fixture
def reporte(tmp_path):
    """Cuatro ciclos: multi-configuración (C&1), diferido (D&1), herencia del mes anterior
    (A&1, con Detalle_Frontera) y exención sin historia (E&1)."""
    detalle = [
        # C&1: dos configuraciones, la más cara (C_2) fija la partida; la detención OT paga por SSCC
        _bloque("C&1", "C", "C_1", T0, tarifa_p=10, tarifa_d=1, fuente=RIO_C_PARTIDA),
        _bloque("C&1", "C", "C_2", T0, tarifa_p=20, tarifa_d=2, fuente=RIO_C_PARTIDA),
        _bloque("C&1", "C", "C_1", T0 + Q, tarifa_p=10, tarifa_d=1, motivo="OT", comentario="Presta SSCC", fuente=RIO_C_DETENCION),
        _bloque("C&1", "C", "C_2", T0 + Q, tarifa_p=20, tarifa_d=2, motivo="OT", comentario="Presta SSCC", fuente=RIO_C_DETENCION),
        # D&1: diferido (sigue el próximo mes)
        _bloque("D&1", "D", "D_1", T0 + pd.Timedelta(days=1)),
        _bloque("D&1", "D", "D_1", T0 + pd.Timedelta(days=1) + Q),
        # A&1: la partida está en julio (frontera); en agosto solo la detención
        _bloque("A&1", "A", "A_1", T0, tarifa_p=5, tarifa_d=1),
        # E&1: sin instrucción y con exención
        _bloque("E&1", "E", "E_1", T0 + pd.Timedelta(days=2), tarifa_p=3, tarifa_d=0, motivo="Sin_Registro_RIO",
                eo="Sin_Registro_RIO", exencion=True),
    ]
    frontera = [_bloque("A&1", "A", "A_1", T0 - Q, tarifa_p=5, tarifa_d=1)]
    ciclos = pd.DataFrame([
        _ciclo("C&1", "C", "Inicia y termina este mes", T0, T0 + Q, 20.0, 2.0, 4.0, motivo_d="OT",
               coment_d="Presta SSCC", fuente_p=RIO_C_PARTIDA, fuente_d=RIO_C_DETENCION),
        _ciclo("D&1", "D", "Continua proximo mes", T0 + pd.Timedelta(days=1), T0 + pd.Timedelta(days=1) + Q,
               0.0, 0.0, 2.0, empresa="E2"),
        _ciclo("A&1", "A", "Viene del mes anterior", T0 - Q, T0, 5.0, 1.0, 2.0),
        _ciclo("E&1", "E", "Inicia y termina este mes", T0 + pd.Timedelta(days=2), T0 + pd.Timedelta(days=2),
               3.0, 0.0, 1.0, motivo_p="Sin_Registro_RIO", motivo_d="Sin_Registro_RIO", exencion=True),
    ])
    rio = pd.DataFrame([
        {"FECHA_HORA_RIO": RIO_C_PARTIDA, "NOMBRE CONFIGURACIÓN": "C_2", "Central_Relacionada_RIO": "C",
         "U. GENERADORA": "C", "CONSIGNAS": "PP", "MOTIVO": "OM", "ESTADO OPERACIONAL": "N", "COMENTARIO": "E/S", "Mes": "anterior"},
        {"FECHA_HORA_RIO": RIO_C_DETENCION, "NOMBRE CONFIGURACIÓN": "C_2", "Central_Relacionada_RIO": "C",
         "U. GENERADORA": "C", "CONSIGNAS": "", "MOTIVO": "OT", "ESTADO OPERACIONAL": "N", "COMENTARIO": "Presta SSCC", "Mes": "actual"},
        {"FECHA_HORA_RIO": T0 + pd.Timedelta(hours=5), "NOMBRE CONFIGURACIÓN": "Z_1", "Central_Relacionada_RIO": "Z",
         "U. GENERADORA": "Z", "CONSIGNAS": "", "MOTIVO": "OM", "ESTADO OPERACIONAL": "N", "COMENTARIO": "nada", "Mes": "actual"},
    ])
    costos = pd.DataFrame([
        {"DIA": 260801, "HORA": 1, "UNIDAD": u, "Partida_Fria": p, "Partida_Tibia": 0, "Partida_Caliente": 0,
         "Detencion": d, "Partida_Tibia_2": 0, "Llave_Concatenada": f"260801-1{u}", "Costo_Cero": "NO",
         "Fria_Num1_M": 24, "Tibia_Num1_O": 0, "Tibia_Num2_N": None, "Caliente_Num1_P": 24}
        for u, p, d in (("C_1", 10, 1), ("C_2", 20, 2), ("D_1", 10, 1), ("A_1", 5, 1), ("E_1", 3, 0))])
    central_empresa = pd.DataFrame({"Central": ["C", "D", "A"], "Empresa": ["E1", "E2", "E1"]})  # E se rescata
    empresas = pd.DataFrame({"Empresa": ["E1", "E2"], "Total_SC_PD_CLP": [18 + 4 + 2, 0]})
    ruta = tmp_path / "Reporte.xlsx"
    with pd.ExcelWriter(ruta) as w:
        ciclos.to_excel(w, sheet_name="Resumen_Ciclos_PD", index=False)
        pd.DataFrame(detalle).to_excel(w, sheet_name="Detalle_15Min", index=False)
        pd.DataFrame(frontera).to_excel(w, sheet_name="Detalle_Frontera", index=False)
        empresas.to_excel(w, sheet_name="SC_por_Empresa", index=False)
        rio.to_excel(w, sheet_name="RIO_Usado", index=False)
        costos.to_excel(w, sheet_name="Costos_PD_Usados", index=False)
        central_empresa.to_excel(w, sheet_name="Central_Empresa", index=False)
        pd.DataFrame({"interruptor": ["RESOLUCION_MARGEN", "MARGEN_NETEADO_POR_CICLO"],
                      "valor": ["bloque", 0]}).to_excel(w, sheet_name="Parametros_Motor", index=False)
    return ruta


def _csv(carpeta, hoja):
    return pd.read_csv(next(carpeta.glob(f"*_{hoja}.csv")), encoding="utf-8-sig")


def test_carpeta_libro_y_menu(reporte):
    carpeta = generar_entrega(reporte, version="Definitivo")
    assert len(list(carpeta.glob("*.csv"))) == 12
    wb = load_workbook(next(carpeta.glob("*.xlsx")))
    assert wb.sheetnames == SHEETS
    assert wb["Menu"]["A2"].value == 260801 and wb["Menu"]["D2"].value == "Definitivo"


def test_encabezados_y_letras_del_excel_horario(reporte):
    carpeta = generar_entrega(reporte)
    wb = load_workbook(next(carpeta.glob("*.xlsx")))

    def fila1(hoja, n):
        return [wb[hoja].cell(row=1, column=i + 1).value or "" for i in range(n)]  # celda vacía = ""
    assert fila1("Sobrecosto_PD xHyC", 35) == XHYC[:35]          # A..AI
    assert fila1("PARTIDAS_DETENCIONES", 30) == PD[:30]          # A..AD
    assert fila1("Sobrecosto_Ciclo", 19) == CICLO[:19]           # A..S
    assert fila1("RESUMEN", 5) == ["Empresa", "PAGA", "RECIBE", "SALDO", "rep"]
    assert wb["Ciclos inconclusos"]["A1"].value == "Próximo mes" and wb["Ciclos inconclusos"]["S1"].value == "Mes anterior"
    assert wb["Costos_de_P-D"]["A1"].value == "id" and wb["Costos_de_P-D"]["Q1"].value == "Costo Partida tibia 2"


def test_llaves_y_marcas_de_partida_detencion(reporte):
    carpeta = generar_entrega(reporte)
    x = _csv(carpeta, "Sobrecosto_PD_xHyC")
    c1 = x[x["Ciclo de operación"] == "C&1"].sort_values(["central", "FECHA_HORA"])
    assert c1.iloc[0]["Id"] == "26080100:00C_1" and c1.iloc[1]["Id"] == "26080100:15C_1"
    # SI exactamente en el primer y último bloque de cada configuración
    assert list(c1["Proceso_Partida"].fillna("")) == ["SI", "", "SI", ""]
    assert list(c1["proceso_detencion"].fillna("")) == ["", "SI", "", "SI"]
    p = _csv(carpeta, "PARTIDAS_DETENCIONES")
    pc = p[p["Clave Ciclo"] == "C&1"].sort_values("FECHA_HORA")
    assert list(pc["Proceso_Partida"].fillna("")) == ["SI", ""] and list(pc["proceso_detencion"].fillna("")) == ["", "SI"]
    assert pc.iloc[0]["Ciclo + fecha + hora"] == "26080100:00C" and int(pc.iloc[0]["Ciclo"]) == 1
    assert list(pc["hora"]) == ["00:00", "00:15"] and list(pc["cuarto"]) == [1, 2]
    # el diferido no tiene detención; el que viene de julio no tiene partida en agosto
    assert (p.loc[p["Clave Ciclo"] == "D&1", "proceso_detencion"].fillna("") == "").all()
    assert (p.loc[p["Clave Ciclo"] == "A&1", "Proceso_Partida"].fillna("") == "").all()
    assert (p.loc[p["Clave Ciclo"] == "A&1", "proceso_detencion"].fillna("") == "SI").any()


def test_valores_reproducen_al_motor(reporte):
    """Lo que las fórmulas darán en Excel, calculado en pandas, iguala Resumen_Ciclos_PD."""
    carpeta = generar_entrega(reporte)
    c = _csv(carpeta, "Sobrecosto_Ciclo").set_index("Ciclo de operación")
    for check in ("Check SC", "Check partida", "Check detención", "Check margen"):
        assert (c[check].abs() <= 1).all(), c[check]
    assert c.loc["C&1", "Total Costos Partida"] == 20 and c.loc["C&1", "Total Costos Detención"] == 2
    assert c.loc["C&1", "Total Margen"] == 4 and c.loc["C&1", "Total Sobrecosto_P-D"] == 18
    # herencia: la partida y el margen de julio entran por F y G
    assert c.loc["A&1", "Total Costos Partida"] == 0 and c.loc["A&1", "Total Costos Partida ciclo inconcluso"] == 5
    assert c.loc["A&1", "Margen ciclo inconcluso"] == 1 and c.loc["A&1", "Total Sobrecosto_P-D"] == 4
    # exención sin historia paga sin instrucción
    assert c.loc["E&1", "Total Costos Partida"] == 3
    # diferido: se traspasa y no entra al RESUMEN
    assert c.loc["D&1", "Ciclo completo"] == 0 and c.loc["D&1", "Empresa"] == "Se traspasa al proximo mes"
    r = _csv(carpeta, "RESUMEN").set_index("Empresa")
    assert r.loc["E1", "RECIBE"] == 24 and r.loc["E2", "RECIBE"] == 0 and (r["CHECK"].abs() <= 1).all()
    # candidatas: la configuración cara gana la partida y SSCC habilita la detención OT
    x = _csv(carpeta, "Sobrecosto_PD_xHyC")
    u = x[(x["Ciclo de operación"] == "C&1") & (x["Proceso_Partida"] == "SI")].set_index("central")["COSTO_PARTIDA [$]"]
    assert u["C_1"] == 10 and u["C_2"] == 20
    p = _csv(carpeta, "PARTIDAS_DETENCIONES")
    det = p[(p["Clave Ciclo"] == "C&1") & (p["proceso_detencion"] == "SI")].iloc[0]
    assert det["Instrucción detención"] == "OT" and det["Presta SSCC detención"] == 1 and det["Monto Detenciones"] == 2


def test_plantillas_de_formulas(reporte):
    todas = " ".join(f for hoja in FORMULAS.values() for f in hoja.values())
    assert "[@" not in todas and "[[#This Row]" not in todas
    assert "$U$2:$U${N}" in FORMULAS["PARTIDAS_DETENCIONES"]["S"]
    assert all("{r}" in f for hoja in FORMULAS.values() for f in hoja.values())
    assert FORMULAS["Sobrecosto_Ciclo"]["H"] == "=IF(C{r}+D{r}+F{r}>E{r}+G{r},C{r}+D{r}+F{r}-E{r}-G{r},0)"


def test_formulas_escritas_en_el_libro(reporte):
    carpeta = generar_entrega(reporte)
    wb = load_workbook(next(carpeta.glob("*.xlsx")))
    x = wb["Sobrecosto_PD xHyC"]
    assert x["U2"].value.startswith('=IFERROR(IF(N2="SI",VLOOKUP(AH2&D2,\'Costos_de_P-D\'!$A$2:$R$')
    assert x["AB2"].value == "=Z2*Y2*AA2" and x["AC2"].value == "=J2" and x["O2"].value == "=+I2"
    p = wb["PARTIDAS_DETENCIONES"]
    filas_partida = [r for r in range(2, p.max_row + 1) if p.cell(row=r, column=6).value == "SI"]
    assert filas_partida
    s = p.cell(row=filas_partida[0], column=19).value
    assert s.startswith("=IFERROR(_xlfn.MAXIFS('Sobrecosto_PD xHyC'!$U$2:$U$") and "AE" in s
    filas_sin = [r for r in range(2, p.max_row + 1) if p.cell(row=r, column=6).value != "SI"]
    assert p.cell(row=filas_sin[0], column=19).value == 0  # valor, no fórmula, fuera de los extremos
    assert p["D2"].value.startswith("=SUMIF('Sobrecosto_PD xHyC'!$AT$2:$AT$")
    c = wb["Sobrecosto_Ciclo"]
    assert c["C2"].value.startswith("=SUMIF(PARTIDAS_DETENCIONES!$N$2:$N$")
    assert c["E2"].value.startswith("=SUMIF('Sobrecosto_PD xHyC'!$AC$2:$AC$")
    assert c["F2"].value.startswith("=SUMIF('Ciclos inconclusos'!$S$3:$S$")
    assert c["H2"].value == "=IF(C2+D2+F2>E2+G2,C2+D2+F2-E2-G2,0)"
    assert c["I2"].value.startswith("=IF(B2=1,VLOOKUP(Q2,Central_Empresa!$A$2:$B$")
    assert c["T2"].value.startswith("=ROUND((H2-") and c["W2"].value.startswith("=ROUND((E2+G2-")
    r = wb["RESUMEN"]
    assert r["C2"].value.startswith("=SUMIF(Sobrecosto_Ciclo!$I$2:$I$") and r["G1"].value == "=SUM(D:D)"
    inc = wb["Ciclos inconclusos"]
    assert inc["S3"].value == "=T3" and "'xHyC mes anterior'" in inc["X3"].value
    assert inc["C3"].value.startswith("=SUMIF(PARTIDAS_DETENCIONES!")


def test_instrucciones_rio_sscc_y_ciclo_que_la_uso(reporte):
    carpeta = generar_entrega(reporte)
    rio = _csv(carpeta, "Instrucciones_RIO")
    por_config = rio.set_index("Configuración")
    assert list(rio["SSCC"]) == [1 if "SSCC" in str(c) else 0 for c in rio["COMENTARIO"]]
    assert (rio.loc[rio["COMENTARIO"] == "Presta SSCC", "Ciclo"] == "C&1").all()
    assert (rio.loc[rio["COMENTARIO"] == "Presta SSCC", "Usada en"] == "Detencion").all()
    assert (rio.loc[rio["Configuración"] == "Z_1", "Ciclo"] == "No encontrado").all()
    vacio = rio.loc[rio["COMENTARIO"] == "E/S", "Clave Ciclo partida"].iloc[0]  # MOTIVO OM, no PP/PMT
    assert pd.isna(vacio) or vacio == ""
    assert pd.isna(por_config.loc["Z_1", "Combustible Partida"]) or por_config.loc["Z_1", "Combustible Partida"] == ""
    c = _csv(carpeta, "Sobrecosto_Ciclo").set_index("Ciclo de operación")
    assert c.loc["C&1", "SSCC"] == 1


def test_central_empresa_incluye_el_rescate_del_motor(reporte):
    carpeta = generar_entrega(reporte)
    ce = _csv(carpeta, "Central_Empresa").set_index("Central")
    assert ce.loc["E", "Empresa"] == "E1" and ce.loc["E", "Origen"].startswith("rescate")
    assert ce.loc["C", "Origen"] == "diccionario"


def test_sin_rio_ni_costos_se_degrada_a_encabezados(reporte, tmp_path):
    hojas = pd.read_excel(reporte, sheet_name=None)
    for h in ("RIO_Usado", "Costos_PD_Usados", "Central_Empresa"):
        hojas.pop(h)
    ruta = tmp_path / "Reporte_viejo.xlsx"
    with pd.ExcelWriter(ruta) as w:
        for h, df in hojas.items():
            df.to_excel(w, sheet_name=h, index=False)
    carpeta = generar_entrega(ruta, tmp_path / "salida")
    wb = load_workbook(next(carpeta.glob("*.xlsx")))
    assert wb["Instrucciones RIO"].max_row == 1 and wb["Costos_de_P-D"].max_row == 1
    c = _csv(carpeta, "Sobrecosto_Ciclo").set_index("Ciclo de operación")
    assert (c["Check SC"].abs() <= 1).all()   # los valores siguen cerrando sin insumos


@pytest.fixture
def retiros(tmp_path):
    """Retiros a 15 min (negativos, como vienen): cuarto 1 = T0, 2 = T0+15, 193 = día 3 00:00."""
    ruta = tmp_path / "Retiros_15min.csv"
    pd.DataFrame({"Cuarto de Hora": [1, 1, 2, 2, 193],
                  "Suministrador": ["S1", "S2", "S1", "E1", "S2"],
                  "Medida_kWh": [-30, -10, -20, -20, -5]}).to_csv(ruta, index=False, sep=";")
    return ruta


def test_prorrateo_llena_paga_y_cuadra_con_recibe(reporte, retiros):
    carpeta = generar_entrega(reporte, retiros=retiros)
    wb = load_workbook(next(carpeta.glob("*.xlsx")))
    assert wb.sheetnames == SHEETS[:SHEETS.index("RESUMEN") + 1] + [HOJA_PAGOS] + SHEETS[SHEETS.index("RESUMEN") + 1:]
    # C&1 (18) cruza cuartos 1-2, A&1 (4) el 1, E&1 (2) el 193; D&1 es diferido y no se reparte
    p = _csv(carpeta, "Cuadro_de_pagos").set_index(["Ciclo de operación", "Suministrador"])
    assert list(p.columns) == PAGOS[2:]
    assert p.loc[("C&1", "S1"), "PAGA"] == pytest.approx(18 * 50 / 80)
    assert p.loc[("C&1", "E1"), "Prorrata"] == pytest.approx(20 / 80)
    assert p.loc[("A&1", "S2"), "PAGA"] == pytest.approx(1)
    assert p.loc[("E&1", "S2"), "PAGA"] == pytest.approx(2)
    assert "D&1" not in p.index.get_level_values(0)
    r = _csv(carpeta, "RESUMEN").set_index("Empresa")
    assert r.loc["S1", "PAGA"] == pytest.approx(14.25) and r.loc["S2", "PAGA"] == pytest.approx(5.25)
    # E1 recibe 24 y además paga como suministrador: el saldo neto es 19,5
    assert r.loc["E1", "PAGA"] == pytest.approx(4.5) and r.loc["E1", "SALDO"] == pytest.approx(19.5)
    assert r["SALDO"].sum() == pytest.approx(0)
    assert (r["CHECK"].abs() <= 1).all()
    res, cuadro = wb["RESUMEN"], wb[HOJA_PAGOS]
    assert res["B2"].value.startswith(f"=SUMIF('{HOJA_PAGOS}'!$B$2:$B$")
    assert cuadro["E2"].value == "=IF(D2=0,0,C2/D2)" and cuadro["G2"].value == "=E2*F2"
    assert cuadro["F2"].value.startswith("=SUMIFS(Sobrecosto_Ciclo!$H$2:$H$")
    assert any("PRORRATEO" in str(c.value) for c in wb["Leeme"]["A"])
    assert (carpeta / "SCPD_2608_Prorrateo_Detalle_15min.csv").exists()
    menu = [c.value for c in wb["Menu"]["A"]]
    assert "Retiros_15min.csv" in menu


def test_sin_retiros_paga_queda_en_cero(reporte):
    carpeta = generar_entrega(reporte)
    r = _csv(carpeta, "RESUMEN")
    assert (r["PAGA"] == 0).all() and not list(carpeta.glob("*Cuadro_de_pagos*"))


def test_sin_retiros_elimina_archivos_de_pagos_de_corrida_anterior(reporte, retiros):
    carpeta = generar_entrega(reporte, retiros=retiros)
    assert (carpeta / "SCPD_2608_Cuadro_de_pagos.csv").exists()
    assert (carpeta / "SCPD_2608_Prorrateo_Detalle_15min.csv").exists()

    assert generar_entrega(reporte) == carpeta
    assert not (carpeta / "SCPD_2608_Cuadro_de_pagos.csv").exists()
    assert not (carpeta / "SCPD_2608_Prorrateo_Detalle_15min.csv").exists()
    assert len(list(carpeta.glob("*.csv"))) == 12
