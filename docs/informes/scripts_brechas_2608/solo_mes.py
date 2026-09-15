"""Agosto 2026: comparacion restringida a ciclos que INICIAN Y TERMINAN dentro del mes.
Excel corregido: se excluyen ciclos traspasados, ciclos con herencia de 'Ciclos inconclusos'
y ciclos cuyo primer bloque con generacion es la hora 1 del 1-ago (vienen de julio).
Motor: Estado_Ciclo_Mes == 'Inicia y termina este mes'."""
import openpyxl, pandas as pd
pd.set_option("display.width", 250); pd.set_option("display.float_format", lambda v: f"{v:,.0f}")
S = r"C:\Users\ANGEL~1.SIL\AppData\Local\Temp\claude\C--Kpi-SC-PD-5-a-6-Claude\993c7f16-aeaf-4559-9bdd-7f17479d2ac6\scratchpad"
NEW = r"C:\Kpi SC_PD\5 a 6\Sobrecostos_PD_2608 pre fixed.xlsm"
SAL = r"C:\Kpi SC_PD\5 a 6\Claude\sc-pd-15min-main\Carpeta_de_Trabajo\Reporte_Sobrecostos_PD_Final.xlsx"

# ---- Excel: Sobrecosto_Ciclo con columnas de herencia ----
wb = openpyxl.load_workbook(NEW, read_only=True, data_only=True, keep_links=False)
ws = wb["Sobrecosto_Ciclo"]; rows = []
for f in ws.iter_rows(min_row=2, values_only=True):
    if f[0] is None: continue
    n = lambda v: float(v) if isinstance(v, (int, float)) else 0.0
    rows.append({"Ciclo": str(f[0]), "Central": str(f[0]).split("&")[0], "Partida": n(f[2]), "Detencion": n(f[3]),
                 "Margen": n(f[4]), "P_inc": n(f[5]), "M_inc": n(f[6]), "SC": n(f[7]),
                 "Empresa": f[8] if isinstance(f[8], str) else "", "Traspasa": isinstance(f[8], str) and f[8].startswith("Se traspasa")})
ex = pd.DataFrame(rows)
# ---- Excel: fechas de cada ciclo desde xHyC (hoja completa) ----
ws = wb["Sobrecosto_PD xHyC"]; it = ws.iter_rows(values_only=True); hdr = list(next(it))
ix = {h: i for i, h in enumerate(hdr) if h is not None}
fechas = []
for r in it:
    if r[ix["Ciclo de operación"]] is None or not r[ix["generacion"]]: continue
    fechas.append((str(r[ix["Ciclo de operación"]]), int(r[ix["fecha"]]), int(r[ix["hora"]]), float(r[ix["generacion"]])))
wb.close()
fx = pd.DataFrame(fechas, columns=["Ciclo", "fecha", "hora", "gen"])
fx["fh"] = pd.to_datetime(fx["fecha"].astype(str), format="%y%m%d") + pd.to_timedelta(fx["hora"] - 1, unit="h")
rng = fx.groupby("Ciclo").agg(inicio=("fh", "min"), fin=("fh", "max"), gen=("gen", "sum"))
ex = ex.merge(rng, left_on="Ciclo", right_index=True, how="left")
ex["hereda"] = (ex.P_inc != 0) | (ex.M_inc != 0)
ex["desde_julio"] = ex["inicio"] <= pd.Timestamp("2026-08-01 00:00")
ex["solo_mes"] = ~ex.Traspasa & ~ex.hereda & ~ex.desde_julio
print(f"Excel: {len(ex)} ciclos | traspasados {int(ex.Traspasa.sum())} | con herencia {int(ex.hereda.sum())} | "
      f"parten en hora 1 {int(ex.desde_julio.sum())} | SOLO MES {int(ex.solo_mes.sum())}")
exm = ex[ex.solo_mes]

# ---- Motor ----
mo = pd.read_excel(SAL, sheet_name="Resumen_Ciclos_PD")
mo["Empresa"] = mo["Empresa"].astype(str).str.strip()
print("Motor Estado_Ciclo_Mes:", mo["Estado_Ciclo_Mes"].value_counts().to_dict())
mom = mo[mo["Estado_Ciclo_Mes"] == "Inicia y termina este mes"]

print(f"\nTOTAL SOLO MES:  Excel {exm.SC.sum():,.0f} ({len(exm)} ciclos)   Motor {mom['Total SC_PD'].sum():,.0f} ({len(mom)} ciclos)   "
      f"delta {mom['Total SC_PD'].sum() - exm.SC.sum():+,.0f} ({100*(mom['Total SC_PD'].sum()/exm.SC.sum()-1):+.1f}%)")
print(f"  Excel: partida {exm.Partida.sum():,.0f}  det {exm.Detencion.sum():,.0f}  margen {exm.Margen.sum():,.0f}")
print(f"  Motor: partida {mom['Costo_Partida_Efectivo'].sum():,.0f}  det {mom['Costo_Detencion_Efectivo'].sum():,.0f}  margen {mom['Margen_Suma_Ciclo'].sum():,.0f}")

e = pd.concat([exm.groupby("Empresa").agg(Ex_ciclos=("SC", "size"), Ex_SC=("SC", "sum")),
               mom.groupby("Empresa").agg(Mo_ciclos=("Total SC_PD", "size"), Mo_SC=("Total SC_PD", "sum"))], axis=1).fillna(0)
e = e[e.index != ""]; e["Delta"] = e.Mo_SC - e.Ex_SC
e["pct"] = (e.Delta / e.Ex_SC.where(e.Ex_SC > 0) * 100)
e = e.sort_values("Delta", key=abs, ascending=False)
print("\nPOR EMPRESA (solo ciclos que inician y terminan en agosto):")
print(e.to_string(formatters={"pct": lambda v: "n/a" if pd.isna(v) else f"{v:+.0f}%"}))
print(f"\n  suma |delta| = {e.Delta.abs().sum():,.0f}")

c = pd.concat([exm.groupby("Central").agg(Ex_ciclos=("SC", "size"), Ex_partida=("Partida", "sum"), Ex_det=("Detencion", "sum"), Ex_margen=("Margen", "sum"), Ex_SC=("SC", "sum")),
               mom.groupby("Central_Relacionada").agg(Mo_ciclos=("Total SC_PD", "size"), Mo_partida=("Costo_Partida_Efectivo", "sum"), Mo_det=("Costo_Detencion_Efectivo", "sum"), Mo_margen=("Margen_Suma_Ciclo", "sum"), Mo_SC=("Total SC_PD", "sum"))], axis=1).fillna(0)
c["Delta"] = c.Mo_SC - c.Ex_SC
c = c[c.Delta.abs() > 100000].sort_values("Delta", key=abs, ascending=False)
print("\nPOR CENTRAL (|delta| > 100 k):"); print(c.to_string())
with pd.ExcelWriter(S + r"\SoloMes_2608.xlsx", engine="xlsxwriter") as w:
    e.reset_index().to_excel(w, sheet_name="Por_Empresa", index=False)
    c.reset_index().to_excel(w, sheet_name="Por_Central", index=False)
    exm.to_excel(w, sheet_name="Excel_ciclos", index=False)
    mom.to_excel(w, sheet_name="Motor_ciclos", index=False)
ex.to_pickle(S + r"\excel_fixed_fechas.pkl")
