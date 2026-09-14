"""Agosto 2026: Excel horario vs salida v7 (tarifa maxima, con empalme julio) para las
empresas que el usuario marco. Por central y por ciclo, con causas."""
from collections import defaultdict

import openpyxl
import pandas as pd

pd.set_option("display.width", 250); pd.set_option("display.float_format", lambda v: f"{v:,.0f}")
pd.set_option("display.max_colwidth", 60)
XLSM = r"C:\Kpi SC_PD\5 a 6\Claude\Sobrecostos_PD_2608 pre.xlsm"
SAL = r"C:\Kpi SC_PD\5 a 6\Claude\sc-pd-15min-main\Carpeta_de_Trabajo\Reporte_Sobrecostos_PD_Final.xlsx"
EMPRESAS = ["ENGIE", "GMETROPOLITANA", "SGA", "BE FORESTALES", "ANTILHUE", "COLMITO",
            "ENLASA", "ELEKTRAGEN", "ENERGIA_SIETE", "NUEVA DEGAN"]

# ---------- Excel: ciclos ----------
wb = openpyxl.load_workbook(XLSM, read_only=True, data_only=True, keep_links=False)
ws = wb["Sobrecosto_Ciclo"]
hdr = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
ex = []
for f in ws.iter_rows(min_row=2, values_only=True):
    if f[0] is None:
        continue
    n = lambda v: float(v) if isinstance(v, (int, float)) else 0.0
    ciclo, completo, c_p, c_d, margen, f_i, g_i, sc, emp = f[:9]
    ex.append({"Ciclo": str(ciclo), "Central": str(ciclo).split("&")[0], "Completo": completo,
               "Partida": n(c_p) + n(f_i), "Detencion": n(c_d), "Margen": n(margen) + n(g_i),
               "SC": n(sc), "Empresa": emp if isinstance(emp, str) else "",
               "Traspasa": isinstance(emp, str) and emp.startswith("Se traspasa")})
wb.close()
ex = pd.DataFrame(ex)
print("Excel columnas:", hdr[:12])

# ---------- motor ----------
mo = pd.read_excel(SAL, sheet_name="Resumen_Ciclos_PD")
det = pd.read_excel(SAL, sheet_name="Detalle_15Min")

for emp in EMPRESAS:
    e = ex[ex["Empresa"].astype(str).str.strip().str.upper() == emp.upper()]
    centrales_ex = set(e["Central"])
    m = mo[mo["Empresa"].astype(str).str.strip().str.upper() == emp.upper()]
    centrales_mo = set(m["Central_Relacionada"])
    # traspasados del excel para esas centrales
    tr = ex[ex["Traspasa"] & ex["Central"].isin(centrales_ex | centrales_mo)]
    print("\n" + "=" * 120)
    print(f"{emp}   Excel {e['SC'].sum():,.0f}   Motor {m['Total SC_PD'].sum():,.0f}   delta {m['Total SC_PD'].sum() - e['SC'].sum():+,.0f}")
    print("=" * 120)
    print("  centrales Excel:", sorted(centrales_ex), "| centrales motor:", sorted(centrales_mo))
    a = e.groupby("Central").agg(Ex_ciclos=("SC", "size"), Ex_partida=("Partida", "sum"), Ex_detencion=("Detencion", "sum"),
                                 Ex_margen=("Margen", "sum"), Ex_SC=("SC", "sum"))
    b = m.groupby("Central_Relacionada").agg(Mo_ciclos=("Total SC_PD", "size"), Mo_partida=("Costo_Partida_Efectivo", "sum"),
                                              Mo_detencion=("Costo_Detencion_Efectivo", "sum"), Mo_margen=("Margen_Suma_Ciclo", "sum"),
                                              Mo_SC=("Total SC_PD", "sum"))
    c = a.join(b, how="outer").fillna(0)
    c["Ex_traspasa"] = c.index.map(lambda x: int((tr["Central"] == x).sum()))
    print(c.to_string())
    # ciclos del motor: observaciones
    print("  -- motor: Obs_Partida --")
    print("    " + m["Obs_Partida"].value_counts().to_string().replace("\n", "\n    "))
    print("  -- motor: Obs_Detencion --")
    print("    " + m["Obs_Detencion"].value_counts().to_string().replace("\n", "\n    "))
    print("  -- motor: Estado_Ciclo_Mes --")
    print("    " + m["Estado_Ciclo_Mes"].value_counts().to_string().replace("\n", "\n    "))
    cols = ["Etiqueta_Relacionada", "Inicio_Ciclo", "Termino_Ciclo", "Horas_Detenida_Ciclo", "Generacion_Suma_Ciclo",
            "Margen_Suma_Ciclo", "Tipo_Partida", "Config_Tarifa_Partida", "Costo_Partida_Efectivo", "Obs_Partida",
            "Costo_Detencion_Efectivo", "Obs_Detencion", "Total SC_PD"]
    print("  -- motor: ciclos --")
    print(m[cols].sort_values("Inicio_Ciclo").to_string(index=False))
    print("  -- Excel: ciclos --")
    print(e[["Ciclo", "Completo", "Partida", "Detencion", "Margen", "SC"]].to_string(index=False))
    if not tr.empty:
        print("  -- Excel: traspasados --")
        print(tr[["Ciclo", "Partida", "Detencion", "Margen"]].to_string(index=False))
